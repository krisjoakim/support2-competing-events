#!/usr/bin/env python3
"""Build project_report.docx from the canonical Markdown; no analysis is run.

Dependencies: pandoc (or pypandoc_binary), python-docx.
Usage: python report/build_report.py [--pandoc /path/to/pandoc]
"""
from pathlib import Path
import argparse
import os
import re
import shutil
import subprocess
import tempfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / 'project_report.md'
OUTPUT = ROOT / 'project_report.docx'


def element(tag, **attrs):
    node = OxmlElement(tag)
    for key, value in attrs.items():
        node.set(qn(key), str(value))
    return node


def style_document(path):
    doc = Document(path)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.78)
    section.footer_distance = Inches(.3)
    width = 6.94
    for name in ['Normal', 'Body Text', 'First Paragraph', 'Compact']:
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor(0, 0, 0)
        pf = style.paragraph_format
        pf.space_after = Pt(7)
        pf.line_spacing = 1.08
        pf.widow_control = True
    for name, size in [('Title', 23), ('Subtitle', 12), ('Heading 1', 15), ('Heading 2', 12)]:
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = name.startswith('Heading')
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if style.font._element.rPr is not None:
            for color in style.font._element.rPr.findall(qn('w:color')):
                for key in list(color.attrib):
                    if 'theme' in key:
                        del color.attrib[key]
        style.paragraph_format.space_before = Pt(12 if name.startswith('Heading') else 0)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True
    for name in ['Caption', 'Image Caption']:
        if name in doc.styles:
            doc.styles[name].font.name = 'Times New Roman'
            doc.styles[name].font.size = Pt(10)
    paragraphs = doc.paragraphs
    references = False
    for i, p in enumerate(paragraphs):
        pf = p.paragraph_format
        if i == 0:
            p.style = 'Title'
        elif i == 1:
            p.style = 'Subtitle'
        elif i == 2:
            for run in p.runs:
                run.font.size = Pt(10)
            pf.space_after = Pt(14)
        elif re.match(r'^\d+\. ', p.text) and p.style.name.startswith('Heading'):
            p.style = 'Heading 1'
            references = p.text.startswith('14.')
            if p.text.startswith(('2.', '14.')):
                pf.page_break_before = True
        if i > 2 and paragraphs[i-1].style.name.startswith('Heading'):
            pf.keep_together = True
        if p.text.endswith(':') or p.text.endswith('the update is') or p.text.endswith('defined as'):
            pf.keep_with_next = True
        if p._p.xpath('./w:pPr/w:numPr') and not references:
            pf.keep_together = True
            if i+1 < len(paragraphs) and paragraphs[i+1]._p.xpath('./w:pPr/w:numPr'):
                pf.keep_with_next = True
        if p._p.xpath('.//w:drawing'):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.keep_with_next = True
            pf.keep_together = True
            pf.space_before = Pt(7)
            pf.space_after = Pt(4)
        if p.text.startswith('Figure '):
            p.style = 'Caption'
            pf.keep_with_next = False
            pf.keep_together = True
            pf.space_after = Pt(10)
        if p._p.xpath('./m:oMathPara'):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.keep_together = True
            pf.space_before = Pt(5)
            pf.space_after = Pt(8)
        if references and not p.style.name.startswith('Heading'):
            pf.keep_together = True
            for run in p.runs:
                run.font.size = Pt(10)
    for shape in doc.inline_shapes:
        ratio = shape.height / shape.width
        image_width = min(width, 4.3 / ratio)
        shape.width, shape.height = Inches(image_width), Inches(image_width * ratio)
    # Intentional column widths, matching the seven source tables in order.
    widths = [[1.1, 1.05, 4.79], [3.4, 1.25, 2.29], [1.35, 1.7, 1.8, 2.09],
              [1.4, 2.8, 2.74], [2.4, 2.4, 2.14], [1.6, 2.9, 2.44], [3.7, 1.7, 1.54]]
    assert len(doc.tables) == len(widths), 'Review table widths after adding/removing source tables.'
    for table, sizes in zip(doc.tables, widths):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        borders = element('w:tblBorders')
        for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            borders.append(element('w:' + edge, **{'w:val': 'single', 'w:sz': 4, 'w:color': 'D9D9D9'}))
        table._tbl.tblPr.append(borders)
        margins = element('w:tblCellMar')
        for edge in ['top', 'bottom', 'left', 'right']:
            margins.append(element('w:' + edge, **{'w:w': 85, 'w:type': 'dxa'}))
        table._tbl.tblPr.append(margins)
        for col, size in zip(table.columns, sizes):
            col.width = Inches(size)
        for row_index, row in enumerate(table.rows):
            props = row._tr.get_or_add_trPr()
            props.append(element('w:cantSplit'))
            if row_index == 0:
                props.append(element('w:tblHeader'))
            for col_index, (cell, size) in enumerate(zip(row.cells, sizes)):
                cell.width = Inches(size)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for old in cell._tc.get_or_add_tcPr().findall(qn('w:tcBorders')):
                    cell._tc.get_or_add_tcPr().remove(old)
                cell_borders = element('w:tcBorders')
                for edge in ['top', 'left', 'bottom', 'right']:
                    cell_borders.append(element('w:' + edge, **{'w:val': 'single', 'w:sz': 4, 'w:color': 'D9D9D9'}))
                cell._tc.get_or_add_tcPr().append(cell_borders)
                if row_index == 0:
                    cell._tc.get_or_add_tcPr().append(element('w:shd', **{'w:fill': 'E8EEF3'}))
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.keep_with_next = row_index < len(table.rows) - 1
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_index == 0 or (len(sizes)==3 and sizes[-1]>4 and col_index==2) else WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.font.size = Pt(10)
                        run.font.bold = row_index == 0
                        run.font.color.rgb = RGBColor(0, 0, 0)
    # Page numbers help navigate a long technical report.
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    field = element('w:fldSimple', **{'w:instr': 'PAGE'})
    footer._p.append(field)
    doc.core_properties.title = paragraphs[0].text
    doc.core_properties.subject = paragraphs[1].text
    doc.core_properties.author = 'Kris-Joakim Buttedal-Fyllingen'
    doc.core_properties.language = 'en-GB'
    doc.core_properties.comments = ''
    doc.save(OUTPUT)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--pandoc', default=os.environ.get('PANDOC') or shutil.which('pandoc'))
    args = parser.parse_args()
    if not args.pandoc:
        try:
            import pypandoc
            args.pandoc = pypandoc.get_pandoc_path()
        except ImportError:
            parser.error('Install Pandoc or pypandoc_binary, or supply --pandoc.')
    source = SOURCE.read_text(encoding='utf-8')
    images = re.findall(r'!\[[^\]]*\]\(([^)]+)\)', source)
    for target in images:
        if not (ROOT / target).is_file():
            raise FileNotFoundError(f'Missing source figure: {target}')
    with tempfile.TemporaryDirectory(prefix='support2-docx-') as temporary:
        raw = Path(temporary) / 'converted.docx'
        subprocess.run([args.pandoc, str(SOURCE), '--from=markdown+tex_math_single_backslash+autolink_bare_uris-implicit_figures',
                        '--to=docx', '--standalone', '--resource-path=' + str(ROOT),
                        '--metadata=lang:en-GB', '--output=' + str(raw)], cwd=ROOT, check=True)
        style_document(raw)
    # Verify package contents; visual review remains a required separate step.
    doc = Document(OUTPUT)
    assert len(doc.inline_shapes) == len(images), 'A source figure was not embedded.'
    assert len(doc._element.xpath('.//m:oMathPara')) == len(re.findall(r'^\\\[$', source, re.M)), 'Display equation conversion failed.'
    assert not any('Figure image not supplied' in p.text for p in doc.paragraphs)
    print(f'Built {OUTPUT}: {len(images)} figures, {len(doc.tables)} tables, native Word equations.')


if __name__ == '__main__':
    main()
